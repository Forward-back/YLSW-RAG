"""
多格式文件解析器。返回统一 ParsedDocument 结构。
"""

import os
from dataclasses import dataclass, field
from typing import List, Dict, Optional

# 模块级 PDF 库可用性缓存
try:
    import fitz
    HAS_FITZ = True
except ImportError:
    HAS_FITZ = False


@dataclass
class TableData:
    """解析出的结构化表格。"""
    table_name: str
    sheet_name: str = ""
    columns: List[str] = field(default_factory=list)
    rows: List[Dict[str, str]] = field(default_factory=list)
    markdown: str = ""


@dataclass
class ParsedDocument:
    """解析器统一输出。"""
    file_name: str
    file_type: str
    full_text: str = ""
    tables: List[TableData] = field(default_factory=list)
    image_descriptions: List[str] = field(default_factory=list)
    parse_error: Optional[str] = None


# ───────────────────────────────────────────────────────
# TXT
# ───────────────────────────────────────────────────────

def parse_txt(file_path: str) -> ParsedDocument:
    doc = ParsedDocument(file_name=os.path.basename(file_path), file_type="txt")
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            doc.full_text = f.read()
    except UnicodeDecodeError:
        try:
            with open(file_path, "r", encoding="gbk") as f:
                doc.full_text = f.read()
        except Exception as e:
            doc.parse_error = f"TXT 解析失败: {e}"
    except Exception as e:
        doc.parse_error = f"TXT 解析失败: {e}"
    return doc


# ───────────────────────────────────────────────────────
# MD（已清洗 Markdown 直通，保留标题/表格结构原样）
# ───────────────────────────────────────────────────────

def parse_md(file_path: str) -> ParsedDocument:
    doc = ParsedDocument(file_name=os.path.basename(file_path), file_type="md")
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            doc.full_text = f.read()
    except UnicodeDecodeError:
        try:
            with open(file_path, "r", encoding="gbk") as f:
                doc.full_text = f.read()
        except Exception as e:
            doc.parse_error = f"MD 解析失败: {e}"
    except Exception as e:
        doc.parse_error = f"MD 解析失败: {e}"
    return doc


# ───────────────────────────────────────────────────────
# PDF
# ───────────────────────────────────────────────────────

def parse_pdf(file_path: str) -> ParsedDocument:
    doc = ParsedDocument(file_name=os.path.basename(file_path), file_type="pdf")
    try:
        if HAS_FITZ:
            text_parts = []
            with fitz.open(file_path) as pdf:
                for page in pdf:
                    page_text = page.get_text()
                    if page_text.strip():
                        text_parts.append(page_text)
            doc.full_text = "\n".join(text_parts)
        else:
            raise ImportError("fitz not available")
    except ImportError:
        try:
            import pdfplumber
            text_parts = []
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    t = page.extract_text()
                    if t:
                        text_parts.append(t)
            doc.full_text = "\n".join(text_parts)
        except Exception as e:
            doc.parse_error = f"PDF 解析失败: {e}"
    except Exception as e:
        doc.parse_error = f"PDF 解析失败: {e}"

    if len(doc.full_text.strip()) < 100 and not doc.parse_error:
        doc.parse_error = "PDF 文本量过少，需要 PaddleOCR 处理（扫描件）"
    return doc


# ───────────────────────────────────────────────────────
# DOCX
# ───────────────────────────────────────────────────────

def parse_docx(file_path: str) -> ParsedDocument:
    doc = ParsedDocument(file_name=os.path.basename(file_path), file_type="docx")
    try:
        from docx import Document
        from docx.text.paragraph import Paragraph
        d = Document(file_path)

        text_parts = []
        t_idx = 0

        for element in d.element.body:
            tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag
            if tag == "p":
                para = Paragraph(element, d)
                if para.text.strip():
                    text_parts.append(para.text.strip())
            elif tag == "tbl":
                if t_idx < len(d.tables):
                    table = d.tables[t_idx]
                    t_idx += 1
                    rows_data = []
                    columns = []
                    for i, row in enumerate(table.rows):
                        cells = [cell.text.strip() for cell in row.cells]
                        if i == 0:
                            columns = cells
                        else:
                            row_dict = {columns[j]: cells[j] for j in range(min(len(columns), len(cells)))}
                            rows_data.append(row_dict)

                    if rows_data:
                        table_name = f"{os.path.basename(file_path)}_表{t_idx}"
                        md_rows = ["| " + " | ".join(columns) + " |",
                                   "| " + " | ".join(["---"] * len(columns)) + " |"]
                        for row in rows_data:
                            md_rows.append("| " + " | ".join(row.get(c, "") for c in columns) + " |")
                        md_table = "\n".join(md_rows)

                        doc.tables.append(TableData(
                            table_name=table_name,
                            columns=columns,
                            rows=rows_data,
                            markdown=md_table,
                        ))

        doc.full_text = "\n".join(text_parts)
    except Exception as e:
        doc.parse_error = f"DOCX 解析失败: {e}"
    return doc


# ───────────────────────────────────────────────────────
# XLSX
# ───────────────────────────────────────────────────────

def parse_xlsx(file_path: str) -> ParsedDocument:
    doc = ParsedDocument(file_name=os.path.basename(file_path), file_type="xlsx")
    try:
        from openpyxl import load_workbook
        wb = load_workbook(file_path, data_only=True, read_only=True)

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows_raw = list(ws.iter_rows(values_only=True))
            if not rows_raw:
                continue

            headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(rows_raw[0])]
            rows_data = []
            for row in rows_raw[1:]:
                if all(v is None for v in row):
                    continue
                row_dict = {headers[j]: str(row[j]) if row[j] is not None else ""
                            for j in range(len(headers))}
                rows_data.append(row_dict)

            if rows_data:
                table_name = os.path.splitext(os.path.basename(file_path))[0]
                if len(wb.sheetnames) > 1:
                    table_name += f"_{sheet_name}"

                md_rows = ["| " + " | ".join(headers) + " |",
                           "| " + " | ".join(["---"] * len(headers)) + " |"]
                for row in rows_data:
                    md_rows.append("| " + " | ".join(row.get(h, "") for h in headers) + " |")
                md_table = "\n".join(md_rows)

                doc.tables.append(TableData(
                    table_name=table_name,
                    sheet_name=sheet_name,
                    columns=headers,
                    rows=rows_data,
                    markdown=md_table,
                ))

                if "COMMENTS" in headers and "COLUMN_NAME" in headers:
                    desc = (f"本表为数据表字段字典，共 {len(rows_data)} 个字段。"
                            f"字段：{'、'.join(headers)}。每个字段的业务含义以 COMMENTS 列为准。")
                else:
                    desc = (f"本表为结构化数据表，共 {len(rows_data)} 行 {len(headers)} 列。"
                            f"字段：{'、'.join(headers)}。具体数据内容可通过数据库查询获取。")
                doc.full_text += f"\n\n## {table_name}\n\n{desc}\n"

        wb.close()
    except Exception as e:
        doc.parse_error = f"XLSX 解析失败: {e}"
    return doc


# ───────────────────────────────────────────────────────
# DOC (旧格式，antiword → LibreOffice → Word COM 多级 fallback)
# ───────────────────────────────────────────────────────

def _doc_via_utf16_scan(file_path: str) -> str | None:
    """纯 Python 兜底：扫描 WordDocument 流中的 UTF-16LE 文本段。

    用于 antiword 无法解析的 .doc（如 WPS 生成的文档缺少标准 piece table）。
    按偏移拼接正文段，过滤域代码（EMBED 等）与二进制噪声段。
    """
    try:
        import olefile
        import struct
    except ImportError:
        return None
    try:
        ole = olefile.OleFileIO(file_path)
        if not ole.exists("WordDocument"):
            return None
        wd = ole.openstream("WordDocument").read()
    except Exception:
        return None

    flags = struct.unpack_from("<H", wd, 0x0A)[0]
    if not flags & 0x1000:  # fExtChar 未设置时为 CP1252 文本，不做启发式提取
        return None

    min_chars = 60  # 域代码/噪声段普遍 <60 字符，正文段更长
    segs = []
    i = 0
    n = len(wd)
    while i + 2 <= n:
        try:
            ch = wd[i:i + 2].decode("utf-16-le")
        except UnicodeDecodeError:
            i += 2
            continue
        if not (ch.isprintable() or ch in "\r\n\t"):
            i += 2
            continue
        start = i
        while i + 2 <= n:
            try:
                ch = wd[i:i + 2].decode("utf-16-le")
            except UnicodeDecodeError:
                break
            if not (ch.isprintable() or ch in "\r\n\t"):
                break
            i += 2
        if i - start >= min_chars * 2:
            segs.append(wd[start:i].decode("utf-16-le"))
        i += 2

    kept = [s for s in segs if _segment_quality(s) >= 0.9]
    if not kept:
        return None
    return "\n".join(kept)


def _segment_quality(text: str) -> float:
    """正文段中 CJK 汉字 + ASCII 可打印 + 换行的占比（用于过滤二进制噪声）。"""
    good = sum(
        1 for ch in text
        if "一" <= ch <= "鿿" or ch.isascii() or ch in "\r\n\t"
    )
    return good / len(text)


def _doc_via_antiword(file_path: str) -> str | None:
    import subprocess
    result = subprocess.run(
        ["antiword", file_path], capture_output=True, text=True, timeout=30
    )
    if result.returncode == 0 and result.stdout.strip():
        return result.stdout
    return None


def _doc_via_libreoffice(file_path: str) -> str | None:
    """用 LibreOffice headless 将 .doc 转为纯文本后读取。"""
    import subprocess
    import tempfile
    import shutil

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    out_dir = tempfile.mkdtemp(prefix="doc_convert_")
    try:
        result = subprocess.run(
            [soffice, "--headless", "--convert-to", "txt:Text",
             "--outdir", out_dir, file_path],
            capture_output=True, text=True, timeout=120,
        )
        if result.returncode != 0:
            return None
        base = os.path.splitext(os.path.basename(file_path))[0]
        txt_path = os.path.join(out_dir, base + ".txt")
        if not os.path.exists(txt_path):
            return None
        with open(txt_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
    except Exception:
        return None
    finally:
        shutil.rmtree(out_dir, ignore_errors=True)


def _doc_via_word_com(file_path: str) -> str | None:
    """用本机 MS Word COM 自动化转换（Windows + 已安装 Office）。"""
    try:
        import win32com.client
    except ImportError:
        return None
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        try:
            doc = word.Documents.Open(file_path, ReadOnly=True)
            text = doc.Content.Text
            doc.Close(False)
            return text
        finally:
            word.Quit()
    except Exception:
        return None


def parse_doc(file_path: str) -> ParsedDocument:
    doc = ParsedDocument(file_name=os.path.basename(file_path), file_type="doc")

    try:
        text = _doc_via_antiword(file_path)
    except FileNotFoundError:
        text = None
    except Exception:
        text = None

    if text is None:
        text = _doc_via_utf16_scan(file_path)
    if text is None:
        text = _doc_via_libreoffice(file_path)
    if text is None:
        text = _doc_via_word_com(file_path)

    if text is None:
        doc.parse_error = (
            "DOC 解析失败：需要安装 antiword / LibreOffice / MS Word 之一"
        )
    else:
        doc.full_text = text
    return doc


# ───────────────────────────────────────────────────────
# 调度器
# ───────────────────────────────────────────────────────

PARSER_MAP = {
    "txt": parse_txt,
    "md": parse_md,
    "pdf": parse_pdf,
    "docx": parse_docx,
    "xlsx": parse_xlsx,
    "doc": parse_doc,
}


def parse_file(file_path: str) -> ParsedDocument:
    """根据扩展名路由到对应解析器。"""
    ext = os.path.splitext(file_path)[1].lower().lstrip(".")
    parser = PARSER_MAP.get(ext)
    if parser is None:
        doc = ParsedDocument(file_name=os.path.basename(file_path), file_type=ext)
        doc.parse_error = f"不支持的文件类型: .{ext}"
        return doc
    return parser(file_path)
